import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from pptx import Presentation
from pptx.util import Pt
from openai import OpenAI
import re
import threading
import os

# 配置 DeepSeek API

# 改为从环境变量读取
import os
client = OpenAI(
    api_key=os.getenv("DEEPSEEK_API_KEY", ""),
    base_url="https://api.deepseek.com/v1"
)

# 在导入部分添加
from tkinter import simpledialog

class PPTTranslatorApp:
    def __init__(self, root):
        self.root = root
        self.root.title("PPT文档翻译工具")
        self.root.geometry("800x600")
        
        # 语言选项
        self.languages = {
            "中文": "zh",
            "英文": "en",
            "日文": "ja",
            "韩文": "ko",
            "法文": "fr",
            "德文": "de",
            "西班牙文": "es",
            "葡萄牙文": "pt",
            "俄文": "ru",
            "马来文": "ms",
            "泰文": "th"
        }
        
        # 添加翻译引擎配置
        self.translation_engines = {
            "DeepSeek": {"base_url": "https://api.deepseek.com/v1", "model": "deepseek-chat"},
            "OpenAI": {"base_url": "https://api.openai.com/v1", "model": "gpt-3.5-turbo"},
            "Google": {"base_url": "", "model": "gemini-pro"},
            "百度翻译": {"base_url": "", "model": ""},
            "本地Ollama": {"base_url": "http://localhost:11434", "model": ""}
        }
        self.current_engine = "DeepSeek"
        self.client = None
        self.setup_client()
        
        # 创建UI
        self.create_widgets()
        
        # 进度条相关
        self.progress = 0
        self.total_slides = 0
        self.current_slide = 0
        
    def create_widgets(self):
        # 文件选择部分
        file_frame = ttk.LabelFrame(self.root, text="PPT文件选择", padding=10)
        file_frame.pack(fill=tk.X, padx=10, pady=5)
        
        self.file_path = tk.StringVar()
        ttk.Entry(file_frame, textvariable=self.file_path, width=50).grid(row=0, column=0, padx=5)
        ttk.Button(file_frame, text="选择文件", command=self.select_file).grid(row=0, column=1, padx=5)
        
        # 语言选择部分
        lang_frame = ttk.LabelFrame(self.root, text="翻译语言设置", padding=10)
        lang_frame.pack(fill=tk.X, padx=10, pady=5)
        
        # 源语言
        ttk.Label(lang_frame, text="源语言:").grid(row=0, column=0, padx=5, pady=5)
        self.source_lang = tk.StringVar(value="中文")
        ttk.Combobox(lang_frame, textvariable=self.source_lang, 
                    values=list(self.languages.keys()), state="readonly").grid(row=0, column=1, padx=5, pady=5)
        
        # 目标语言
        ttk.Label(lang_frame, text="目标语言:").grid(row=0, column=2, padx=5, pady=5)
        self.target_lang = tk.StringVar(value="英文")
        ttk.Combobox(lang_frame, textvariable=self.target_lang, 
                    values=list(self.languages.keys()), state="readonly").grid(row=0, column=3, padx=5, pady=5)
        
        # 输出设置
        output_frame = ttk.LabelFrame(self.root, text="输出设置", padding=10)
        output_frame.pack(fill=tk.X, padx=10, pady=5)
        
        self.output_path = tk.StringVar()
        ttk.Entry(output_frame, textvariable=self.output_path, width=50).grid(row=0, column=0, padx=5)
        ttk.Button(output_frame, text="另存为", command=self.select_output).grid(row=0, column=1, padx=5)
        
        # 进度显示
        self.progress_frame = ttk.LabelFrame(self.root, text="转换进度", padding=10)
        self.progress_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        self.progress_bar = ttk.Progressbar(self.progress_frame, orient=tk.HORIZONTAL, length=600, mode='determinate')
        self.progress_bar.pack(pady=10)
        
        self.progress_label = ttk.Label(self.progress_frame, text="准备就绪")
        self.progress_label.pack()
        
        self.detail_text = tk.Text(self.progress_frame, height=15, state=tk.DISABLED)
        self.detail_text.pack(fill=tk.BOTH, expand=True, pady=5)
        
        # 操作按钮
        button_frame = ttk.Frame(self.root)
        button_frame.pack(fill=tk.X, padx=10, pady=10)
        
        ttk.Button(button_frame, text="开始转换", command=self.start_translation).pack(side=tk.LEFT, padx=5)
        # 在操作按钮部分添加引擎设置按钮
        ttk.Button(button_frame, text="翻译引擎设置", command=self.show_engine_settings).pack(side=tk.LEFT, padx=5)
        ttk.Button(button_frame, text="退出", command=self.root.quit).pack(side=tk.RIGHT, padx=5)
    
    def select_file(self):
        file_path = filedialog.askopenfilename(
            title="选择PPT文件",
            filetypes=[("PPT文件", "*.pptx *.ppt"), ("所有文件", "*.*")]
        )
        if file_path:
            self.file_path.set(file_path)
            # 自动设置输出路径
            dir_name, file_name = os.path.split(file_path)
            name, ext = os.path.splitext(file_name)
            self.output_path.set(os.path.join(dir_name, f"{name}_translated{ext}"))
    
    def select_output(self):
        output_path = filedialog.asksaveasfilename(
            title="保存翻译后的PPT",
            defaultextension=".pptx",
            filetypes=[("PPT文件", "*.pptx"), ("所有文件", "*.*")]
        )
        if output_path:
            self.output_path.set(output_path)
    
    def update_progress(self, current, total, message):
        self.current_slide = current
        self.total_slides = total
        self.progress = int((current / total) * 100) if total > 0 else 0
        
        self.progress_bar["value"] = self.progress
        self.progress_label.config(text=f"进度: {self.progress}% ({current}/{total}页)")
        
        self.detail_text.config(state=tk.NORMAL)
        self.detail_text.insert(tk.END, message + "\n")
        self.detail_text.see(tk.END)
        self.detail_text.config(state=tk.DISABLED)
        
        self.root.update_idletasks()
    
    def process_shape(self, shape):
        """递归处理所有形状（含组合形状嵌套）"""
        # 处理文本框
        if shape.has_text_frame:
            # 收集所有文本段落
            all_text = []
            for paragraph in shape.text_frame.paragraphs:
                paragraph_text = ""
                for run in paragraph.runs:
                    paragraph_text += run.text
                if paragraph_text.strip():
                    all_text.append(paragraph_text)
            
            # 合并所有文本并翻译
            if all_text:
                combined_text = "\n".join(all_text)
                translated = self.translate_text(combined_text)
                
                # 替换第一个run的文本，删除其他run
                if shape.text_frame.paragraphs:
                    first_paragraph = shape.text_frame.paragraphs[0]
                    if first_paragraph.runs:
                        first_run = first_paragraph.runs[0]
                        first_run.text = translated
                        first_run.font.size = Pt(20)  # 设置字号为20磅
                        first_run.font.name = "黑体"  # 设置字体为黑体
                        
                        # 删除其他run
                        for paragraph in shape.text_frame.paragraphs:
                            while len(paragraph.runs) > 1:
                                r = paragraph.runs[1]
                                paragraph._p.remove(r._r)
        
        # 处理表格
        if shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    if cell.text and cell.text.strip():  # 更严格的空值检查
                        translated = self.translate_text(cell.text)
                        if translated:  # 确保翻译结果不为空
                            cell.text = translated
                        for paragraph in cell.text_frame.paragraphs:
                            for run in paragraph.runs:
                                run.font.size = Pt(20)  # 设置字号为20磅
                                run.font.name = "黑体"  # 设置字体为黑体
    
        # 处理组合形状（GroupShape）
        if shape.shape_type == 6:  # GroupShape 的类型标识为6
            for sub_shape in shape.shapes:
                self.process_shape(sub_shape)  # 递归处理子形状
    
    def translate_ppt(self):
        """主处理函数"""
        input_path = self.file_path.get()
        output_path = self.output_path.get()
        
        if not input_path:
            messagebox.showerror("错误", "请选择要翻译的PPT文件")
            return
        
        if not output_path:
            messagebox.showerror("错误", "请设置输出文件路径")
            return
        
        try:
            prs = Presentation(input_path)
            total_slides = len(prs.slides)
            
            for i, slide in enumerate(prs.slides, 1):
                self.update_progress(i, total_slides, f"正在处理第 {i} 页...")
                
                for shape in slide.shapes:
                    self.process_shape(shape)  # 处理所有形状
            
            prs.save(output_path)
            self.update_progress(total_slides, total_slides, f"翻译完成，结果已保存至: {output_path}")
            messagebox.showinfo("完成", "转换完毕")
            
        except Exception as e:
            self.update_progress(0, 0, f"处理失败: {e}")
            messagebox.showerror("错误", f"处理失败: {e}")
    
    def start_translation(self):
        """启动翻译线程"""
        if not self.file_path.get():
            messagebox.showerror("错误", "请选择要翻译的PPT文件")
            return
            
        if hasattr(self, '_translation_thread') and self._translation_thread.is_alive():
            messagebox.showwarning("警告", "翻译任务已在运行中")
            return
            
        # 清空进度信息
        self.detail_text.config(state=tk.NORMAL)
        self.detail_text.delete(1.0, tk.END)
        self.detail_text.config(state=tk.DISABLED)
        
        # 在新线程中执行翻译
        thread = threading.Thread(target=self.translate_ppt)
        thread.daemon = True
        thread.start()
        self._translation_thread = thread  # 保存线程引用

    def setup_client(self):
        """初始化翻译客户端"""
        engine_config = self.translation_engines[self.current_engine]
        self.client = OpenAI(
            api_key=self.api_key if hasattr(self, 'api_key') else "",
            base_url=engine_config["base_url"]
        )

    def show_engine_settings(self):
        """显示翻译引擎设置窗口"""
        settings_window = tk.Toplevel(self.root)
        settings_window.title("翻译引擎设置")
        settings_window.geometry("500x400")
        
        # 引擎选择
        engine_frame = ttk.LabelFrame(settings_window, text="选择翻译引擎", padding=10)
        engine_frame.pack(fill=tk.X, padx=10, pady=5)
        
        self.engine_var = tk.StringVar(value=self.current_engine)
        engine_combobox = ttk.Combobox(engine_frame, textvariable=self.engine_var, 
                                     values=list(self.translation_engines.keys()), state="readonly")
        engine_combobox.pack(fill=tk.X)
        
        # API设置
        api_frame = ttk.LabelFrame(settings_window, text="API设置", padding=10)
        api_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=5)
        
        # API Key
        ttk.Label(api_frame, text="API Key:").grid(row=0, column=0, sticky=tk.W, padx=5, pady=5)
        self.api_key_var = tk.StringVar(value=getattr(self, 'api_key', ''))
        ttk.Entry(api_frame, textvariable=self.api_key_var, show="*").grid(row=0, column=1, sticky=tk.EW, padx=5, pady=5)
        
        # Base URL
        ttk.Label(api_frame, text="Base URL:").grid(row=1, column=0, sticky=tk.W, padx=5, pady=5)
        self.base_url_var = tk.StringVar(value=self.translation_engines[self.current_engine]["base_url"])
        ttk.Entry(api_frame, textvariable=self.base_url_var).grid(row=1, column=1, sticky=tk.EW, padx=5, pady=5)
        
        # 模型选择
        ttk.Label(api_frame, text="模型名称:").grid(row=2, column=0, sticky=tk.W, padx=5, pady=5)
        self.model_var = tk.StringVar(value=self.translation_engines[self.current_engine]["model"])
        ttk.Entry(api_frame, textvariable=self.model_var).grid(row=2, column=1, sticky=tk.EW, padx=5, pady=5)
        
        # 保存按钮
        ttk.Button(settings_window, text="保存设置", command=lambda: self.save_engine_settings(
            self.engine_var.get(),
            self.api_key_var.get(),
            self.base_url_var.get(),
            self.model_var.get(),
            settings_window
        )).pack(side=tk.BOTTOM, pady=10)
        
        # 更新UI当引擎改变时
        def on_engine_change(*args):
            engine = self.engine_var.get()
            self.base_url_var.set(self.translation_engines[engine]["base_url"])
            self.model_var.set(self.translation_engines[engine]["model"])
        
        self.engine_var.trace_add("write", on_engine_change)

    def save_engine_settings(self, engine, api_key, base_url, model, window):
        """保存引擎设置"""
        self.current_engine = engine
        self.api_key = api_key
        self.translation_engines[engine]["base_url"] = base_url
        self.translation_engines[engine]["model"] = model
        self.setup_client()
        window.destroy()
        messagebox.showinfo("成功", "翻译引擎设置已保存")

    def translate_text(self, text):
        """调用翻译引擎翻译文本"""
        if not text.strip():
            return text
            
        try:
            source_lang = self.languages[self.source_lang.get()]
            target_lang = self.languages[self.target_lang.get()]
            
            response = self.client.chat.completions.create(
                model=self.translation_engines[self.current_engine]["model"],
                messages=[
                    {"role": "system", "content": f"你是一个专业翻译引擎，仅将{source_lang}翻译成{target_lang}，不要做任何其它解释，保留所有数字、符号、格式和已有内容。"},
                    {"role": "user", "content": text}
                ],
                temperature=0.1
            )
            return response.choices[0].message.content.strip()
        except Exception as e:
            self.update_progress(self.current_slide, self.total_slides, f"翻译失败: {e}")
            return text


if __name__ == "__main__":
    root = tk.Tk()
    app = PPTTranslatorApp(root)
    root.mainloop()