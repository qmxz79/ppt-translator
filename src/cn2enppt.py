from pptx import Presentation
from pptx.util import Pt
from openai import OpenAI
import re
from concurrent.futures import ThreadPoolExecutor
import requests

# 默认配置 DeepSeek API
deepseek_client = OpenAI(
    api_key="sk-22b8fd0b8d1c480cb30328623f4f847a",
    base_url="https://api.deepseek.com/v1"
)

translation_engine = "DeepSeek"
model_option = ""
client = deepseek_client

def set_translation_engine(engine, model):
    global translation_engine, model_option, client
    translation_engine = engine
    model_option = model

    if translation_engine == "DeepSeek":
        client = OpenAI(
            api_key="sk-22b8fd0b8d1c480cb30328623f4f847a",
            base_url="https://api.deepseek.com/v1"
        )
    elif translation_engine == "DomesticAPI":
        # 配置国内大模型API客户端
        client = None  # 替换为实际的国内大模型API客户端
    elif translation_engine == "LocalOllama":
        # 配置本地ollama部署客户端
        client = None  # 本地ollama部署不需要客户端配置
    else:
        raise ValueError("未知的翻译引擎")

def is_chinese(text):
    """检测文本是否包含中文字符"""
    return re.search(r'[\u4e00-\u9fff]', text)

def translate_text(text):
    """调用相应的翻译引擎翻译中文到英文"""
    try:
        if translation_engine == "DeepSeek":
            response = client.chat.completions.create(
                model=model_option,
                messages=[
                    {"role": "system", "content": "你是一个医学专业翻译引擎，**仅翻译中文到英文**，不要做任何其它解释，保留所有数字、符号、格式和英文内容。"},
                    {"role": "user", "content": text}
                ],
                temperature=0.1
            )
            translated_text = response.choices[0].message.content.strip()
        elif translation_engine == "DomesticAPI":
            # 调用国内大模型API的翻译函数
            translated_text = call_domestic_api(text, model_option)  # 替换为实际调用代码
        elif translation_engine == "LocalOllama":
            # 调用本地ollama部署的翻译函数
            translated_text = call_local_ollama(text, model_option)
        else:
            raise ValueError("未知的翻译引擎")

        return translated_text
    except Exception as e:
        print(f"翻译失败: {e}")
        return text  # 失败时返回原文

def call_domestic_api(text, model):
    """调用国内大模型API的翻译函数"""
    # 替换为实际调用代码
    return "翻译后的文本"

def call_local_ollama(text, model):
    """调用本地ollama部署的翻译函数"""
    url = "http://127.0.0.1:11434/api/chat"  # 替换为实际的 Ollama API URL
    payload = {
        "model": model,
        "text": text
    }
    headers = {
        "Content-Type": "application/json"
    }
    response = requests.post(url, json=payload, headers=headers)
    if response.status_code == 200:
        return response.json().get("translated_text", text)
    else:
        raise Exception(f"Ollama API 调用失败: {response.status_code}")

def process_shape(shape):
    """处理单个形状"""
    if shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if is_chinese(run.text):
                    translated = translate_text(run.text)
                    run.text = translated
                    run.font.size = Pt(20)
    elif shape.has_table:
        for row in shape.table.rows:
            for cell in row.cells:
                if is_chinese(cell.text):
                    translated = translate_text(cell.text)
                    cell.text = translated
                    for paragraph in cell.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(20)

def process_slide(slide):
    """处理单个幻灯片"""
    for shape in slide.shapes:
        process_shape(shape)

def translate_ppt(input_path, output_path, update_progress):
    """主处理函数"""
    prs = Presentation(input_path)
    total_slides = len(prs.slides)
    processed_slides = 0

    with ThreadPoolExecutor() as executor:
        futures = [executor.submit(process_slide, slide) for slide in prs.slides]
        for future in futures:
            future.result()
            processed_slides += 1
            update_progress((processed_slides / total_slides) * 100)

    prs.save(output_path)
    print(f"翻译完成，结果已保存至: {output_path}")